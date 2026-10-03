from __future__ import annotations
import csv
import html
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import List

from .models import Segment, TextTrack

_TIME_RE = re.compile(r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})(?:[,.](\d{1,3}))?")
_RANGE_RE = re.compile(
    r"(?P<a>(?:\d{1,2}:)?\d{1,2}:\d{2}(?:[,.]\d{1,3})?)\s*(?:-->|-|–|—|~|至)\s*"
    r"(?P<b>(?:\d{1,2}:)?\d{1,2}:\d{2}(?:[,.]\d{1,3})?)"
)
_POINT_RE = re.compile(r"\[((?:\d{1,2}:)?\d{1,2}:\d{2}(?:[,.]\d{1,3})?)\]")


_SPEAKER_PREFIX_RE = re.compile(r"^\s*((?:spk|speaker)\s*[:#-]?\s*\d+|(?:講者|說話者)\s*\d+)\s*[：:]\s*(.+)$", re.I)

def split_speaker_prefix(text: str, speaker: str = "") -> tuple[str, str]:
    payload = (text or "").strip()
    sp = (speaker or "").strip()
    m = _SPEAKER_PREFIX_RE.match(payload)
    if m:
        if not sp:
            sp = m.group(1).strip()
        payload = m.group(2).strip()
    return sp, payload


def parse_time(value: str) -> float:
    s = str(value).strip()
    if not s:
        return 0.0
    try:
        if re.fullmatch(r"\d+(?:\.\d+)?", s):
            return float(s)
    except Exception:
        pass
    m = _TIME_RE.search(s)
    if not m:
        return 0.0
    hh = int(m.group(1) or 0)
    mm = int(m.group(2) or 0)
    ss = int(m.group(3) or 0)
    ms_raw = m.group(4) or "0"
    ms = int(ms_raw.ljust(3, "0")[:3])
    return hh * 3600 + mm * 60 + ss + ms / 1000.0


def fmt_time(seconds: float, comma: bool = False) -> str:
    seconds = max(0.0, float(seconds))
    hh = int(seconds // 3600)
    mm = int((seconds % 3600) // 60)
    ss = int(seconds % 60)
    ms = int(round((seconds - int(seconds)) * 1000))
    sep = "," if comma else "."
    return f"{hh:02d}:{mm:02d}:{ss:02d}{sep}{ms:03d}"


def split_sentences(text: str) -> List[str]:
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return []
    parts = re.split(r"(?<=[。！？!?；;\.])\s*|\n+", text)
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) <= 1 and len(text) > 52:
        # Smaller fallback chunks improve KTV positioning for punctuation-poor transcripts.
        step = 36
        return [text[i:i+step].strip() for i in range(0, len(text), step) if text[i:i+step].strip()]
    return parts


def parse_srt(text: str, name: str = "SRT") -> TextTrack:
    blocks = re.split(r"\r?\n\s*\r?\n", text.strip())
    segs = []
    for block in blocks:
        lines = [x.strip("\ufeff") for x in block.splitlines() if x.strip()]
        if not lines:
            continue
        time_idx = next((i for i, x in enumerate(lines) if "-->" in x), None)
        if time_idx is None:
            continue
        m = _RANGE_RE.search(lines[time_idx])
        if not m:
            continue
        payload = " ".join(lines[time_idx+1:]).strip()
        speaker = ""
        sm = re.match(r"^([^：:]{1,30})[：:]\s*(.+)$", payload)
        if sm and any(k in sm.group(1).lower() for k in ["speaker", "講者", "說話者", "主持", "老師", "同學"]):
            speaker, payload = sm.group(1).strip(), sm.group(2).strip()
        segs.append(Segment(parse_time(m.group("a")), parse_time(m.group("b")), payload, speaker))
    return TextTrack(name=name, segments=segs, format_name="SRT", timed=bool(segs), estimated=False)


def parse_vtt(text: str, name: str = "VTT") -> TextTrack:
    cleaned = re.sub(r"^WEBVTT.*?(?:\r?\n){2}", "", text.strip(), flags=re.S | re.I)
    tr = parse_srt(cleaned, name=name)
    tr.format_name = "VTT"
    return tr


def parse_inline_point_timeline(text: str, name: str, fmt: str = "TXT") -> TextTrack | None:
    """Parse exports using point markers like [00:00:30] text.

    Metadata before the first marker is ignored. Text between markers is split into
    smaller sentence cues and its interval is interpolated so KTV highlighting can
    follow the recording more closely.
    """
    normalized = re.sub(r"[ \t]+", " ", (text or "").replace("\u00a0", " "))
    markers = list(_POINT_RE.finditer(normalized))
    if not markers:
        return None
    segs: list[Segment] = []
    for i, m in enumerate(markers):
        start = parse_time(m.group(1))
        body_start = m.end()
        body_end = markers[i + 1].start() if i + 1 < len(markers) else len(normalized)
        body = normalized[body_start:body_end].strip()
        body = re.sub(r"^(逐字稿|辨識引擎[：:]?[^\n]{0,40})\s*", "", body, flags=re.I).strip()
        if not body:
            continue
        end = parse_time(markers[i + 1].group(1)) if i + 1 < len(markers) else start + 8.0
        end = max(start + 0.8, end)
        parts = split_sentences(body) or [body]
        weights = [max(1, len(re.sub(r"\s+", "", x))) for x in parts]
        total = max(1, sum(weights))
        cur = start
        for j, part0 in enumerate(parts):
            speaker, part = split_speaker_prefix(part0)
            if not part:
                continue
            nxt = end if j == len(parts) - 1 else cur + (end - start) * (weights[j] / total)
            nxt = max(cur + 0.35, nxt)
            segs.append(Segment(cur, nxt, part, speaker, True))
            cur = nxt
    if not segs:
        return None
    return TextTrack(name=name, segments=segs, format_name=fmt, timed=True, estimated=True)


def parse_txt(text: str, name: str = "TXT") -> TextTrack:
    inline = parse_inline_point_timeline(text, name, "TXT")
    if inline:
        return inline
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    timed = []
    untimed = []
    for line in lines:
        m = _RANGE_RE.search(line)
        if m:
            payload = (line[:m.start()] + " " + line[m.end():]).strip(" []｜|:-–—")
            speaker = ""
            sm = re.match(r"^([^：:]{1,30})[：:]\s*(.+)$", payload)
            if sm:
                speaker, payload = sm.group(1).strip(), sm.group(2).strip()
            timed.append(Segment(parse_time(m.group("a")), parse_time(m.group("b")), payload, speaker))
        else:
            sm = re.match(r"^\[(?P<t>(?:\d{1,2}:)?\d{1,2}:\d{2}(?:[,.]\d{1,3})?)\]\s*(?P<p>.*)$", line)
            if sm:
                start = parse_time(sm.group("t"))
                timed.append(Segment(start, start, sm.group("p").strip()))
            else:
                untimed.append(line)
    if timed:
        for i, seg in enumerate(timed):
            if seg.end <= seg.start:
                seg.end = timed[i+1].start if i + 1 < len(timed) else seg.start + 3.0
        return TextTrack(name=name, segments=timed, format_name="TXT", timed=True, estimated=False)
    full = "\n".join(untimed) if untimed else text
    return TextTrack(name=name, segments=[Segment(0, 0, x) for x in split_sentences(full)], format_name="TXT", timed=False, estimated=True)


def parse_json_text(text: str, name: str = "JSON") -> TextTrack:
    obj = json.loads(text)
    if isinstance(obj, dict) and "tracks" in obj and obj["tracks"]:
        t = obj["tracks"][0]
        return TextTrack.from_dict(t)
    items = obj.get("segments", obj) if isinstance(obj, dict) else obj
    segs = []
    if isinstance(items, list):
        for it in items:
            if not isinstance(it, dict):
                continue
            start = float(it.get("start", it.get("start_time", 0.0)) or 0.0)
            end = float(it.get("end", it.get("end_time", start)) or start)
            txt = str(it.get("text", it.get("content", "")) or "").strip()
            spk = str(it.get("speaker", it.get("speaker_label", "")) or "")
            if txt:
                segs.append(Segment(start, end, txt, spk, bool(it.get("estimated", False))))
    timed = bool(segs) and any(s.end > s.start for s in segs)
    return TextTrack(name=name, segments=segs, format_name="JSON", timed=timed, estimated=not timed)


def parse_csv_file(path: Path) -> TextTrack:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    segs = []
    for r in rows:
        low = {str(k).strip().lower(): v for k, v in r.items()}
        txt = low.get("text") or low.get("content") or low.get("文字") or low.get("逐字稿") or ""
        start = low.get("start") or low.get("開始") or low.get("start_time") or "0"
        end = low.get("end") or low.get("結束") or low.get("end_time") or start
        spk = low.get("speaker") or low.get("講者") or ""
        if str(txt).strip():
            segs.append(Segment(parse_time(start), parse_time(end), str(txt).strip(), str(spk).strip()))
    timed = any(s.end > s.start for s in segs)
    return TextTrack(path.stem, segs, str(path), "CSV", timed, not timed)


def parse_docx(path: Path) -> TextTrack:
    from docx import Document
    doc = Document(str(path))
    segs = []
    # Prefer table-based timeline (Junba exports and many meeting transcripts).
    for table in doc.tables:
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            if not cells:
                continue
            m = _RANGE_RE.search(cells[0])
            if not m:
                continue
            speaker = cells[1] if len(cells) >= 3 else ""
            txt = cells[2] if len(cells) >= 3 else (cells[1] if len(cells) >= 2 else "")
            speaker, txt = split_speaker_prefix(txt, speaker)
            if txt:
                segs.append(Segment(parse_time(m.group("a")), parse_time(m.group("b")), txt, speaker))
    if segs:
        return TextTrack(path.stem, segs, str(path), "DOCX", True, False)
    paras = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    tr = parse_inline_point_timeline(paras, path.stem, "DOCX") or parse_txt(paras, path.stem)
    tr.source_path = str(path)
    tr.format_name = "DOCX"
    return tr


class _Stripper(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, d):
        if d.strip():
            self.parts.append(d.strip())


def parse_html(path: Path) -> TextTrack:
    text = path.read_text(encoding="utf-8", errors="ignore")
    m = re.search(r'<script[^>]+id=["\']junba-ktv-data["\'][^>]*>(.*?)</script>', text, re.S | re.I)
    if m:
        obj = json.loads(html.unescape(m.group(1)))
        tracks = obj.get("tracks", [])
        if tracks:
            tr = TextTrack.from_dict(tracks[0])
            tr.source_path = str(path)
            tr.format_name = "HTML"
            return tr
    s = _Stripper(); s.feed(text)
    tr = parse_txt("\n".join(s.parts), path.stem)
    tr.source_path = str(path); tr.format_name = "HTML"
    return tr


def load_track(path_str: str) -> TextTrack:
    path = Path(path_str)
    ext = path.suffix.lower()
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".doc":
        raise ValueError("目前支援 .docx，不支援舊式 .doc；請先在 Word 另存成 .docx。")
    if ext == ".csv":
        return parse_csv_file(path)
    if ext in {".html", ".htm"}:
        return parse_html(path)
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if ext == ".srt": tr = parse_srt(text, path.stem)
    elif ext == ".vtt": tr = parse_vtt(text, path.stem)
    elif ext == ".json": tr = parse_json_text(text, path.stem)
    else: tr = parse_txt(text, path.stem)
    tr.source_path = str(path)
    return tr


def to_vtt(track: TextTrack) -> str:
    out = ["WEBVTT", ""]
    for i, s in enumerate(track.segments, 1):
        out += [str(i), f"{fmt_time(s.start)} --> {fmt_time(s.end)}", (f"{s.speaker}：" if s.speaker else "") + s.text, ""]
    return "\n".join(out)


def to_srt(track: TextTrack) -> str:
    out = []
    for i, s in enumerate(track.segments, 1):
        out += [str(i), f"{fmt_time(s.start, True)} --> {fmt_time(s.end, True)}", (f"{s.speaker}：" if s.speaker else "") + s.text, ""]
    return "\n".join(out)
